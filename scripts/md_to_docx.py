import sys
import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def markdown_to_docx(md_path, docx_path, logo_path=None, diagrams_dir=None):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()
    
    # Page setup (Standard A4 with 1 inch margin)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Diagram mapping
    diagram_mapping = [
        ("architecture.png", "Gambar 1.5 - Arsitektur Terpadu Ekosistem KostKu"),
        ("context_diagram.png", "Gambar 3.2.1 - Diagram Konteks Sistem KostKu"),
        ("dfd_level1.png", "Gambar 3.2.2 - Data Flow Diagram (DFD) Level 1 KostKu"),
        ("dfd_level2_kamar.png", "Gambar 3.2.2.1 - DFD Level 2 (Manajemen Kamar & Properti)"),
        ("dfd_level2_transaksi.png", "Gambar 3.2.2.2 - DFD Level 2 (Transaksi Sewa & Mitra Afiliasi)"),
        ("erd.png", "Gambar 3.3.1 - Entity Relationship Diagram (ERD) Basis Data KostKu"),
        ("traceability_flow.png", "Gambar 3.6 - Diagram Alur Kerunutan Kebutuhan Sistem")
    ]
    mermaid_counter = 0

    i = 0
    in_code_block = False
    code_lines = []
    is_mermaid = False

    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Code block handling
        if line.startswith('```'):
            if in_code_block:
                # End code block
                if is_mermaid:
                    # Embed diagram image instead of raw code
                    if diagrams_dir and mermaid_counter < len(diagram_mapping):
                        img_filename, caption_text = diagram_mapping[mermaid_counter]
                        img_path = os.path.join(diagrams_dir, img_filename)
                        if os.path.exists(img_path):
                            p_img = doc.add_paragraph()
                            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p_img.paragraph_format.space_before = Pt(8)
                            p_img.paragraph_format.space_after = Pt(4)
                            run_img = p_img.add_run()
                            try:
                                run_img.add_picture(img_path, width=Inches(6.2))
                            except Exception as e:
                                print(f"Could not embed picture {img_path}: {e}")
                            
                            # Caption
                            p_cap = doc.add_paragraph()
                            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p_cap.paragraph_format.space_before = Pt(2)
                            p_cap.paragraph_format.space_after = Pt(12)
                            run_cap = p_cap.add_run(caption_text)
                            run_cap.font.name = 'Arial'
                            run_cap.font.size = Pt(9)
                            run_cap.font.italic = True
                            run_cap.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue 600
                        else:
                            # Fallback if image missing
                            p = doc.add_paragraph()
                            r = p.add_run(f"[{caption_text}]")
                            r.font.italic = True
                    mermaid_counter += 1
                else:
                    # Non-mermaid code block: format in elegant box
                    table = doc.add_table(rows=1, cols=1)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(table, color="CBD5E1", sz="4")
                    cell = table.cell(0, 0)
                    set_cell_shading(cell, "F8FAFC")
                    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    code_text = '\n'.join(code_lines)
                    run = p.add_run(code_text)
                    run.font.name = 'Courier New'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

                code_lines = []
                in_code_block = False
                is_mermaid = False
            else:
                in_code_block = True
                code_lines = []
                is_mermaid = 'mermaid' in line
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^\s*[-*_]{3,}\s*$', line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Markdown Table Detection
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            parsed_rows = []
            for t_line in table_lines:
                cleaned = t_line.strip('|')
                cells = [c.strip() for c in cleaned.split('|')]
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                parsed_rows.append(cells)

            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                for r in parsed_rows:
                    while len(r) < num_cols:
                        r.append('')
                
                table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(table, color="CBD5E1", sz="4")
                
                # Check if it's the UDINUS cover header table (with Logo UDINUS)
                is_cover_table = any('Logo UDINUS' in c for r in parsed_rows for c in r)

                for r_idx, r_data in enumerate(parsed_rows):
                    row = table.rows[r_idx]
                    is_header = (r_idx == 0) and not is_cover_table
                    
                    for c_idx, cell_value in enumerate(r_data):
                        cell = row.cells[c_idx]
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        
                        # Background colors: Taste Skill palette
                        if is_cover_table:
                            set_cell_shading(cell, "F8FAFC")
                        elif is_header:
                            set_cell_shading(cell, "1E3A8A") # Royal Navy Header
                        elif r_idx % 2 == 1:
                            set_cell_shading(cell, "FFFFFF")
                        else:
                            set_cell_shading(cell, "F8FAFC") # Soft light zebra
                        
                        cell_p = cell.paragraphs[0]
                        cell_p.paragraph_format.space_before = Pt(2)
                        cell_p.paragraph_format.space_after = Pt(2)
                        
                        # Logo image in cover table
                        if 'Logo UDINUS' in cell_value and logo_path and os.path.exists(logo_path):
                            run = cell_p.add_run()
                            try:
                                run.add_picture(logo_path, width=Inches(0.9))
                            except Exception:
                                run.text = "[Logo UDINUS]"
                        else:
                            clean_text = cell_value.replace('**', '').replace('`', '').replace('*', '')
                            
                            # Check status color highlights
                            is_added = any(w in clean_text for w in ['Ditambahkan', 'Added'])
                            is_modified = any(w in clean_text for w in ['Diubah', 'Modified'])
                            is_removed = any(w in clean_text for w in ['Dihapus', 'Dieliminasi'])
                            
                            run = cell_p.add_run(clean_text)
                            run.font.name = 'Arial'
                            
                            if is_header:
                                run.font.size = Pt(9.5)
                                run.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # White on Navy
                            elif is_cover_table:
                                run.font.size = Pt(9)
                                run.bold = ('Program Studi' in clean_text or 'Nomor Dokumen' in clean_text or 'Revisi' in clean_text)
                                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                            else:
                                run.font.size = Pt(9.5)
                                if is_added:
                                    run.bold = True
                                    run.font.color.rgb = RGBColor(0x15, 0x80, 0x3D) # Emerald Green
                                elif is_modified:
                                    run.bold = True
                                    run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09) # Amber Gold
                                elif is_removed:
                                    run.bold = True
                                    run.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C) # Rose Red
                                else:
                                    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(0)
                p_spacer.paragraph_format.space_after = Pt(4)
            continue

        # Headings with Taste Skill Colors
        m_h1 = re.match(r'^#\s+(.+)$', line)
        m_h2 = re.match(r'^##\s+(.+)$', line)
        m_h3 = re.match(r'^###\s+(.+)$', line)
        m_h4 = re.match(r'^####\s+(.+)$', line)

        if m_h1:
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(m_h1.group(1).replace('**', ''))
            run.font.name = 'Arial'
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Navy #1E3A8A
            i += 1
            continue
        elif m_h2:
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(m_h2.group(1).replace('**', ''))
            run.font.name = 'Arial'
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Royal Blue #2563EB
            i += 1
            continue
        elif m_h3:
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(m_h3.group(1).replace('**', ''))
            run.font.name = 'Arial'
            run.font.size = Pt(11.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0D, 0x94, 0x88) # Modern Teal #0D9488
            i += 1
            continue
        elif m_h4:
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run = h.add_run(m_h4.group(1).replace('**', ''))
            run.font.name = 'Arial'
            run.font.size = Pt(10.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69) # Slate 600
            i += 1
            continue

        # Bullet list
        m_bullet = re.match(r'^\s*[\*\-]\s+(.+)$', line)
        if m_bullet:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text = m_bullet.group(1)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    # If it's a key term, give subtle navy color
                    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                else:
                    r = p.add_run(part.replace('`', ''))
                    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                r.font.name = 'Arial'
                r.font.size = Pt(10)
            i += 1
            continue

        # Numbered list
        m_num = re.match(r'^\s*(\d+)\.\s+(.+)$', line)
        if m_num:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            num_str = m_num.group(1) + ". "
            r_num = p.add_run(num_str)
            r_num.bold = True
            r_num.font.name = 'Arial'
            r_num.font.size = Pt(10)
            r_num.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue number
            text = m_num.group(2)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                else:
                    r = p.add_run(part.replace('`', ''))
                    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                r.font.name = 'Arial'
                r.font.size = Pt(10)
            i += 1
            continue

        # Cover text block formatting (before tables)
        if i < 15 and ('SKPL' in line or 'SPESIFIKASI' in line or 'KostKu' in line or 'Dipersiapkan oleh' in line or 'Oscar' in line or 'Yohanes' in line or 'Mirza' in line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if 'SKPL' == line.strip():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line)
                run.bold = True
                run.font.size = Pt(18)
                run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            elif 'SPESIFIKASI KEBUTUHAN PERANGKAT LUNAK' in line:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line)
                run.bold = True
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            elif 'KostKu' in line:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line)
                run.bold = True
                run.font.size = Pt(13)
                run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            elif 'untuk:' in line or 'Dipersiapkan oleh:' in line:
                p.paragraph_format.space_before = Pt(6)
                run = p.add_run(line)
                run.bold = True
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            else:
                run = p.add_run(line)
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            else:
                r = p.add_run(part.replace('`', ''))
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            r.font.name = 'Arial'
            r.font.size = Pt(10)
        
        i += 1

    doc.save(docx_path)
    print(f"Generated colorful Taste-Skill DOCX at: {docx_path}")

if __name__ == '__main__':
    md_file = sys.argv[1] if len(sys.argv) > 1 else 'docs/SKPL_KostKu.md'
    out_docx = sys.argv[2] if len(sys.argv) > 2 else '/home/rena/Desktop/SKPL_KostKu.docx'
    logo = 'docs/assets/udinus_logo.jpeg'
    diagrams = 'docs/assets/diagrams'
    markdown_to_docx(md_file, out_docx, logo, diagrams)
