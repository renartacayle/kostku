import os
import json
import glob
from pptx import Presentation
from pptx.util import Inches

def build_pptx():
    json_path = 'scripts/slides_data.json'
    with open(json_path, 'r', encoding='utf-8') as f:
        slides_data = json.load(f)
    
    prs = Presentation()
    # 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    import re
    images = sorted(glob.glob('/tmp/slide_pptx_page-*.png'), key=lambda p: int(re.search(r'-(\d+)\.png$', p).group(1)) if re.search(r'-(\d+)\.png$', p) else 0)
    print(f"Found {len(images)} slide images and {len(slides_data)} slide data records.")
    
    for i, slide_info in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        
        # Add high-res slide background image
        if i < len(images):
            img_path = images[i]
            slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
        
        # Add presenter notes
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        
        script_text = slide_info.get('script', '').strip()
        tips_list = slide_info.get('tips', [])
        
        notes_content = f"【 SLIDE {slide_info.get('num', str(i+1))} 】 {slide_info.get('title', '')}\n"
        notes_content += f"Tag: {slide_info.get('tag', '')}\n\n"
        notes_content += "--- NASKAH PRESENTASI (APA YANG HARUS DIUCAPKAN) ---\n"
        notes_content += script_text + "\n\n"
        
        if tips_list:
            notes_content += "--- TIPS GESTUR & TEKNIS SIDANG ---\n"
            for tip in tips_list:
                notes_content += f"• {tip}\n"
        
        tf.text = notes_content
    
    out_dir = '/home/rena/Downloads'
    out_path = os.path.join(out_dir, 'KostKu_Slide_Presentasi_Sidang_RPL.pptx')
    prs.save(out_path)
    print(f"PPTX saved to: {out_path} ({os.path.getsize(out_path) / 1024 / 1024:.2f} MB)")
    
    # Also copy to artifacts & public
    brain_path = '/home/rena/.gemini/antigravity/brain/6318663c-dc01-4644-8421-dc82f19bcbdc/KostKu_Slide_Presentasi_Sidang_RPL.pptx'
    prs.save(brain_path)
    print(f"PPTX copied to brain artifact: {brain_path}")
    
    public_path = 'public/KostKu_Slide_Presentasi_Sidang_RPL.pptx'
    prs.save(public_path)
    print(f"PPTX copied to public: {public_path}")

if __name__ == '__main__':
    build_pptx()
